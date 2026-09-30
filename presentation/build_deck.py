# -*- coding: utf-8 -*-
"""SIH 2026 idea deck for PS SIH26063 (NCPOR Polar Knowledge Index), Team Anarch.

Built on the official SIH 2026 template with the same visual system as the team's RegimeCast
deck (Canva version): team badge, tagline, icon flow, numbered workflow, prototype screenshot
with callouts, stack table, pipeline, stat icons, risk table, viability cycle, benefit bars,
comparison and reference tables.

Every number is taken from the running prototype. Evaluation scores are placeholders
("[__%]") until the hand-verified test set has been run: they must not be invented.
"""
import math
import os
import re

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
ICONS = os.path.join(HERE, "icons", "png")
SHOTS = os.path.join(HERE, "shots")
SRC = os.path.join(HERE, "template.pptx")
OUT = os.path.join(HERE, "SIH26063_Polar_Knowledge_Index.pptx")

IDEA = "Polar Knowledge Index"
TEAM_FOOTER = "@SIH Idea Team - Anarch"
EVAL = "[__%]"  # filled in after the hand-verified evaluation run

NAVY = RGBColor(0x1E, 0x27, 0x61)
BLUE = RGBColor(0x1F, 0x6F, 0xD1)
ORANGE = RGBColor(0xD9, 0x59, 0x26)
GREEN = RGBColor(0x14, 0x8A, 0x5B)
RED = RGBColor(0xC0, 0x39, 0x2B)
TEAL = RGBColor(0x12, 0x8C, 0x7E)
PURPLE = RGBColor(0x5B, 0x3F, 0xA8)
DARK = RGBColor(0x1F, 0x29, 0x37)
MUTED = RGBColor(0x5B, 0x64, 0x75)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
T_BLUE = RGBColor(0xEA, 0xF2, 0xFC)
T_ORANGE = RGBColor(0xFD, 0xF0, 0xE9)
T_GREEN = RGBColor(0xE8, 0xF6, 0xEF)
T_GREY = RGBColor(0xF3, 0xF4, 0xF7)
T_PEACH = RGBColor(0xFB, 0xE3, 0xC8)
BORDER = RGBColor(0x9D, 0xBC, 0xE8)
BODY = "Calibri"
SYMBOL = "Segoe UI Symbol"

prs = Presentation(SRC)
S = prs.slides

# ------------------------------------------------------------------ helpers (from the RegimeCast build)
TOKEN = re.compile(r"(\[\[.+?\]\]|\{\{.+?\}\}|\(\(.+?\)\)|\*\*.+?\*\*|@@.+?@@)")


def add_rich(p, text, size, color=DARK, bold=False, italic=False, font=BODY):
    """[[blue bold]] {{orange bold}} ((green bold)) **bold** @@blue underlined label@@"""
    for part in TOKEN.split(text):
        if not part:
            continue
        c, b, u = color, bold, False
        if part.startswith("[["):
            part, c, b = part[2:-2], BLUE, True
        elif part.startswith("{{"):
            part, c, b = part[2:-2], ORANGE, True
        elif part.startswith("(("):
            part, c, b = part[2:-2], GREEN, True
        elif part.startswith("**"):
            part, b = part[2:-2], True
        elif part.startswith("@@"):
            part, c, b, u = part[2:-2], BLUE, True, True
        r = p.add_run()
        r.text = part
        f = r.font
        f.size, f.name, f.bold, f.italic = Pt(size), font, b, italic
        f.underline = u
        f.color.rgb = c


def set_bullet(p, char="•", marL=0.17, color=None, font="Arial"):
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(int(Inches(marL))))
    pPr.set("indent", str(int(-Inches(marL))))
    if color is not None:
        clr = pPr.makeelement(qn("a:buClr"), {})
        clr.append(clr.makeelement(qn("a:srgbClr"), {"val": str(color)}))
        pPr.append(clr)
    pPr.append(pPr.makeelement(qn("a:buFont"), {"typeface": font}))
    pPr.append(pPr.makeelement(qn("a:buChar"), {"char": char}))


def textbox(slide, x, y, w, h, paras, size=12, color=DARK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
            bullet=None, after=2, bold=False, italic=False, font=BODY, line=None, bullet_font="Arial",
            bullet_color=None):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = Inches(0.03)
    tf.margin_top = tf.margin_bottom = Inches(0.01)
    tf.vertical_anchor = anchor
    for i, text in enumerate(paras if isinstance(paras, list) else [paras]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if line:
            p.line_spacing = line
        p.space_after = Pt(after)
        if bullet:
            set_bullet(p, bullet, color=bullet_color or BLUE, font=bullet_font)
        add_rich(p, text, size, color, bold, italic, font)
    return tb


def card(slide, x, y, w, h, fill, line=None, radius=0.08, width=1.0):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.adjustments[0] = min(0.5, radius / min(w, h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(width)
    s.shadow.inherit = False
    return s


def shape(slide, kind, x, y, w, h, fill, line=None):
    s = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
    s.shadow.inherit = False
    return s


def icon(slide, name, x, y, size):
    return slide.shapes.add_picture(os.path.join(ICONS, f"{name}.png"), Inches(x), Inches(y), Inches(size), Inches(size))


def icon_circle(slide, x, y, d, name, fill):
    c = shape(slide, MSO_SHAPE.OVAL, x, y, d, d, fill)
    s = d * 0.56
    icon(slide, name, x + (d - s) / 2, y + (d - s) / 2, s)
    return c


def arrow(slide, x1, y1, x2, y2, color=NAVY, width=1.75):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    ln = c.line._get_or_add_ln()
    ln.append(ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"}))
    return c


def line_(slide, x1, y1, x2, y2, color=BORDER, width=2.0):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    return c


def picture(slide, name, x, y, w=None, h=None, border=None):
    kw = {}
    if w is not None:
        kw["width"] = Inches(w)
    if h is not None:
        kw["height"] = Inches(h)
    pic = slide.shapes.add_picture(os.path.join(SHOTS, name), Inches(x), Inches(y), **kw)
    if border is not None:
        pic.line.color.rgb = border
        pic.line.width = Pt(1)
    return pic


def remove(shp):
    el = shp._element
    el.getparent().remove(el)


def shape_named(slide, name):
    return next(s for s in slide.shapes if s.name == name)


def set_title(slide, text, size=None):
    t = next(s for s in slide.shapes if s.is_placeholder and s.placeholder_format.type in (1, 3))
    runs = [r for p in t.text_frame.paragraphs for r in p.runs]
    runs[-1].text = text
    for r in runs[:-1]:
        r.text = ""
    if size:
        runs[-1].font.size = Pt(size)


def team_badge(slide):
    """The template's 'Your Team Name' oval becomes the TEAM ANARCH badge, as in the Canva deck."""
    for shp in slide.shapes:
        if shp.has_text_frame and "Team" in shp.text_frame.text and "Name" in shp.text_frame.text:
            tf = shp.text_frame
            for p in list(tf.paragraphs)[1:]:
                p._p.getparent().remove(p._p)
            p = tf.paragraphs[0]
            for r in list(p.runs):
                r._r.getparent().remove(r._r)
            p.alignment = PP_ALIGN.CENTER
            tf.margin_left = tf.margin_right = 0
            add_rich(p, "TEAM", 13, NAVY, bold=True, font="Arial")
            p2 = tf.add_paragraph()
            p2.alignment = PP_ALIGN.CENTER
            add_rich(p2, "ANARCH", 13, NAVY, bold=True, font="Arial")
            return


def footer(slide):
    for shp in slide.shapes:
        if shp.name.startswith("Footer Placeholder"):
            shp.text_frame.paragraphs[0].runs[0].text = TEAM_FOOTER


def prepare(slide):
    for shp in list(slide.shapes):
        if shp.name == "TextBox 8":
            remove(shp)
    team_badge(slide)
    footer(slide)


def table(slide, x, y, col_w, row_h, rows, header_fill=NAVY, font_size=10, highlight=None,
          first_col_bold=True, align_center_from=1):
    n_rows, n_cols = len(rows), len(rows[0])
    shp = slide.shapes.add_table(n_rows, n_cols, Inches(x), Inches(y), Inches(sum(col_w)), Inches(row_h * n_rows))
    tbl = shp.table
    for i, w in enumerate(col_w):
        tbl.columns[i].width = Inches(w)
    for r in range(n_rows):
        tbl.rows[r].height = Inches(row_h)
        for c in range(n_cols):
            cell = tbl.cell(r, c)
            cell.margin_left = cell.margin_right = Inches(0.05)
            cell.margin_top = cell.margin_bottom = Inches(0.02)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = header_fill if r == 0 else (T_GREY if r % 2 == 0 else WHITE)
            p = cell.text_frame.paragraphs[0]
            cell.text_frame.word_wrap = True
            p.alignment = PP_ALIGN.LEFT if c < align_center_from else PP_ALIGN.CENTER
            color = WHITE if r == 0 else DARK
            bold = r == 0 or (c == 0 and first_col_bold)
            if highlight and (r, c) in highlight:
                color, bold = highlight[(r, c)], True
            add_rich(p, rows[r][c], font_size, color, bold)
    return shp


def section(slide, x, y, w, h, title, items, color=BLUE, icon_name=None, size=10.5, bullet="❖", after=4, fill=WHITE):
    card(slide, x, y, w, h, fill, line=BORDER, radius=0.06, width=1.0)
    tx = x + 0.12
    if icon_name:
        icon_circle(slide, x + 0.1, y + 0.08, 0.32, icon_name, color)
        tx = x + 0.5
    textbox(slide, tx, y + 0.07, w - (tx - x) - 0.1, 0.36, title, size=12.5, color=color, bold=True,
            anchor=MSO_ANCHOR.MIDDLE)
    textbox(slide, x + 0.14, y + 0.5, w - 0.26, h - 0.56, items, size=size, bullet=bullet, after=after,
            bullet_font=SYMBOL, bullet_color=color)


def caps_label(slide, x, y, w, text, color=BLUE, size=12.5, icon_name=None):
    if icon_name:
        icon_circle(slide, x, y + 0.01, 0.3, icon_name, color)
        x, w = x + 0.37, w - 0.37
    return textbox(slide, x, y, w, 0.34, text, size=size, color=color, bold=True, anchor=MSO_ANCHOR.MIDDLE)


def ribbon(slide, x, y, w, h, text, fill=NAVY, size=12):
    """Rounded banner with white bold caps text (Canva deck section headers)."""
    card(slide, x, y, w, h, fill, radius=0.12)
    textbox(slide, x, y, w, h, text, size=size, color=WHITE, bold=True, align=PP_ALIGN.CENTER,
            anchor=MSO_ANCHOR.MIDDLE)


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


# ================================================================== slide 1: title page
s1 = S[0]
fields = shape_named(s1, "TextBox 9")
paras = fields.text_frame.paragraphs
values = {
    1: "Problem Statement ID : 26063",
    2: "Problem Statement Title : Integrated Polar Science Outreach, Knowledge Repository and Media "
       "Dissemination Portal",
    3: "Theme : Smart Education",
    4: "PS Category : Software",
    5: "Team ID :",
    6: "Team Name : Anarch",
}
for i, text in values.items():
    if i >= len(paras):
        break
    for r in list(paras[i]._p.findall(qn("a:r"))):
        paras[i]._p.remove(r)
    paras[i].alignment = PP_ALIGN.LEFT
    r = paras[i].add_run()
    r.text = text
    r.font.size, r.font.name, r.font.bold, r.font.color.rgb = Pt(18), "Arial", True, DARK
for shp in list(s1.shapes):  # the template's "TITLE PAGE" label is not part of the submitted title page
    if shp.has_text_frame and shp.text_frame.text.strip().upper() == "TITLE PAGE":
        remove(shp)
notes(s1, "Polar Knowledge Index: NCPOR's existing archive, harvested, linked to expeditions, searchable and cited "
          "to the page, with reviewed lessons and announcements in English and Hindi.")

# ================================================================== slide 2: proposed solution
s2 = S[1]
prepare(s2)
set_title(s2, IDEA, size=40)
ttl = next(x for x in s2.shapes if x.is_placeholder and x.placeholder_format.type in (1, 3))
ttl.top, ttl.height = Inches(0), Inches(1.02)
tag = textbox(s2, 1.9, 1.06, 9.5, 0.36, "", size=16, align=PP_ALIGN.CENTER)
p = tag.text_frame.paragraphs[0]
add_rich(p, "“Every expedition. Every page. ", 16, DARK, bold=True, font="Times New Roman")
add_rich(p, "Every claim cited.", 16, BLUE, bold=True, font="Times New Roman")
add_rich(p, "”", 16, DARK, bold=True, font="Times New Roman")
tb = textbox(s2, 0.35, 1.46, 12.6, 0.34, "", size=13)
p = tb.text_frame.paragraphs[0]
set_bullet(p, "•", color=DARK)
add_rich(p, "Proposed Solution:  ", 13, DARK, bold=True)
add_rich(p, "A knowledge index built on what NCPOR already publishes — harvested, linked to each expedition, "
            "and cited to the page.", 13, BLUE)

# left column: icon flow + three question blocks
flow = [("archive", NAVY), ("scan", BLUE), ("sitemap2", TEAL), ("search", BLUE), ("clipboard", GREEN), ("school", ORANGE)]
for i, (ic, col) in enumerate(flow):
    x = 0.45 + i * 0.95
    icon_circle(s2, x, 1.88, 0.48, ic, col)
    if i < len(flow) - 1:
        arrow(s2, x + 0.52, 2.12, x + 0.9, 2.12, color=NAVY, width=2.25)
textbox(s2, 0.35, 2.4, 5.8, 0.24, "Harvest  →  Read  →  Link  →  Search  →  Review  →  Teach",
        size=9, color=MUTED, italic=True, align=PP_ALIGN.CENTER)

blocks = [
    ("What it does?", [
        "@@Harvests, never re-enters@@ : DSpace expedition reports (OAI-PMH) + NCPOR news feed, polite & cached "
        "— originals stay where they are",
        "@@Reads the scans@@ : PDF text layer or Tesseract OCR, page image kept for “view original page”",
        "@@Links everything@@ : item → expedition · station · topic · year; machine links stay "
        "unconfirmed until a curator confirms",
    ]),
    ("How it answers the PS?", [
        "@@Repository@@ : 32 expedition reports, 765 papers and 837 news posts in one index, one page per expedition",
        "@@Dissemination@@ : cited answers, lesson units and web + social posts, in English and Hindi",
    ]),
    ("What is different about it?", [
        "@@Real NCPOR archive@@ , zero new data entry for staff",
        "@@Checked, not claimed@@ : every sentence fact-checked against its page; unsupported ones block approval",
        "@@Named reviewer@@ + append-only audit log; refuses when the archive has no source",
        "@@Coverage view@@ : shows which expeditions have no public report or story",
    ]),
]
y = 2.72
for head, items in blocks:
    tb = textbox(s2, 0.35, y, 5.6, 0.28, "", size=12.5)
    p = tb.text_frame.paragraphs[0]
    add_rich(p, "➤ ", 12.5, ORANGE, bold=True, font=SYMBOL)
    add_rich(p, head, 12.5, DARK, bold=True)
    body_h = 0.36 * len(items) + 0.12
    textbox(s2, 0.5, y + 0.3, 5.45, body_h, items, size=9.5, bullet="•", bullet_color=BLUE, after=2.5)
    y += 0.3 + body_h + 0.04

# middle column: numbered workflow
textbox(s2, 6.25, 1.83, 2.2, 0.3, "WORKFLOW", size=13, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
steps = [("DSPACE + RSS", "HARVEST", "archive", NAVY), ("TEXT LAYER", "/ OCR", "scan", BLUE),
         ("LINK TO", "EXPEDITION", "sitemap2", TEAL), ("HYBRID", "SEARCH", "search", BLUE),
         ("CITED AI", "DRAFT", "robot", PURPLE), ("FAITHFULNESS", "CHECK", "listcheck", ORANGE),
         ("NAMED", "REVIEWER", "usercheck", GREEN), ("PUBLISH", "EN / HI", "language", RED)]
top, step = 2.2, 0.575
line_(s2, 7.35, top + 0.22, 7.35, top + step * 7 + 0.22, color=BORDER, width=2.5)
for i, (l1, l2, ic, col) in enumerate(steps):
    yy = top + i * step
    num = shape(s2, MSO_SHAPE.OVAL, 7.13, yy, 0.44, 0.44, col)
    textbox(s2, 7.13, yy, 0.44, 0.44, f"{i + 1:02d}", size=10.5, color=WHITE, bold=True, align=PP_ALIGN.CENTER,
            anchor=MSO_ANCHOR.MIDDLE)
    left = i % 2 == 0
    lx = 6.0 if left else 7.66
    textbox(s2, lx, yy - 0.01, 1.08, 0.46, [f"**{l1}**", f"**{l2}**"], size=8, color=DARK, after=0,
            align=PP_ALIGN.RIGHT if left else PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE)
    icon(s2, f"{ic}_navy", 7.68 if left else 6.76, yy + 0.07, 0.3)

# right column: prototype image + callouts
textbox(s2, 8.75, 1.83, 4.2, 0.3, "PROTOTYPE IMAGE", size=13, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
picture(s2, "s2_page_viewer.png", 8.75, 2.18, w=4.2, border=MUTED)
textbox(s2, 8.75, 4.92, 4.2, 0.4, ["Working prototype · real DSpace scan · 9th Expedition glacier paper,",
                                    "printed p. 248 — scan beside its searchable text"],
        size=8.5, color=MUTED, italic=True, align=PP_ALIGN.CENTER, after=0)
callouts = [("Scan :", "the page exactly as NCPOR published it", NAVY),
            ("Citation :", "every answer links back to this page", ORANGE),
            ("Coverage :", "expeditions with no public report", NAVY)]
for i, (a, b, col) in enumerate(callouts):
    yy = 5.3 + i * 0.44
    shp = shape(s2, MSO_SHAPE.PENTAGON, 9.35, yy, 3.2, 0.38, col)
    textbox(s2, 9.45, yy, 2.95, 0.38, f"**{a}** {b}", size=9.5, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
textbox(s2, 8.75, 6.62, 4.2, 0.3, ["**github repo link :**", "**demo youtube video :**"], size=8.5, after=0)
notes(s2, "Problem: NCPOR's material is stored in six disconnected systems and nothing turns it into outreach. "
          "Solution: harvest what NCPOR already publishes, link every item to its expedition, and generate "
          "reviewed, cited content only from linked sources.")

# ================================================================== slide 3: technical approach
s3 = S[2]
prepare(s3)
textbox(s3, 2.0, 1.02, 9.3, 0.3, "Technology stack & system architecture — built and running end-to-end on "
        "NCPOR’s real archive", size=12, color=NAVY, italic=True, align=PP_ALIGN.CENTER)

ribbon(s3, 0.35, 1.42, 4.3, 0.42, "Technology Stack", fill=NAVY, size=14)
stack = [["LAYER", "TOOL"],
         ["Harvest", "httpx · OAI-PMH · lxml · robots.txt-aware, 1 req/s, cached"],
         ["Documents", "PyMuPDF · Tesseract OCR 5.4 · page images"],
         ["Search", "PostgreSQL 16 full-text + pgvector · multilingual-e5-base"],
         ["AI", "Google Gemini — grounded drafting, checker, Hindi"],
         ["Backend", "Python 3.12 · FastAPI · SQLAlchemy 2 · Alembic"],
         ["Frontend", "Next.js · TypeScript · Tailwind · WCAG 2.1 AA"],
         ["Security", "Staff roles · JWT · bcrypt · append-only audit"],
         ["Quality", "25 automated tests (pytest) · Docker Compose"]]
table(s3, 0.35, 1.92, [1.05, 3.25], 0.34, stack, header_fill=BLUE, font_size=9.5, align_center_from=2)
logos = ["python", "fastapi", "postgresql", "sqlalchemy", "pytest", "docker",
         "nextdotjs", "typescript", "react", "tailwind", "gemini", "huggingface"]
for i, name in enumerate(logos):
    col, row = i % 6, i // 6
    icon(s3, f"logo_{name}", 0.5 + col * 0.7, 5.1 + row * 0.62, 0.44)
textbox(s3, 0.35, 6.35, 4.3, 0.5, ["((✓)) Originals never re-hosted: links + OCR text + page images only",
                                    "((✓)) Public is read-only; every staff action is audited"], size=9.5, after=1)

ribbon(s3, 4.95, 1.42, 8.0, 0.42, "PIPELINE", fill=ORANGE, size=15)
stages = [
    (1, "DATA SOURCES", BLUE, "database", [
        "**DSpace at NCAOR** · OAI-PMH · 32 report communities",
        "**NCPOR news RSS** · 837 posts since 2012",
        "**NPDC** · link-only until NCPOR permits"]),
    (2, "READ & LINK", TEAL, "scan", [
        "Text layer first, Tesseract OCR if not",
        "Printed page numbers · page images",
        "Rules → AI fallback (always unconfirmed)"]),
    (3, "CORE ENGINE", PURPLE, "brain", [
        "Hybrid search: keyword + meaning (RRF)",
        "Refusal below relevance threshold",
        "Grounded drafts: answer · lesson · post · Hindi"]),
    (4, "VERIFY", ORANGE, "listcheck", [
        "Each sentence: supported / partial / unsupported",
        "Unsupported sentences block approval",
        "Named reviewer approves · audit trail"]),
    (5, "SERVE", NAVY, "server", [
        "FastAPI · role-based staff API",
        "Public endpoints read-only",
        "Cached AI responses for offline demo"]),
    (6, "DELIVER", GREEN, "display", [
        "Search · Ask · scan beside text",
        "Expedition pages · coverage view",
        "Lessons & stories in English / Hindi"]),
]
for idx, (n, title, col, ic, items) in enumerate(stages):
    cx = 4.95 + (idx % 3) * 2.7
    cy = 2.0 + (idx // 3) * 2.28
    num = shape(s3, MSO_SHAPE.OVAL, cx, cy, 0.52, 0.52, WHITE, line=col)
    num.line.width = Pt(2.5)
    textbox(s3, cx, cy, 0.52, 0.52, str(n), size=18, color=col, bold=True, align=PP_ALIGN.CENTER,
            anchor=MSO_ANCHOR.MIDDLE)
    textbox(s3, cx + 0.58, cy + 0.08, 2.0, 0.36, title, size=12.5, color=col, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    card(s3, cx, cy + 0.62, 2.55, 1.45, T_GREY, radius=0.08)
    icon(s3, f"{ic}_navy", cx + 2.13, cy + 0.68, 0.3)
    textbox(s3, cx + 0.08, cy + 0.7, 2.05, 1.35, items, size=9, bullet="•", bullet_color=col, after=3)
    if idx % 3 < 2:
        arrow(s3, cx + 2.58, cy + 1.3, cx + 2.68, cy + 1.3, color=col, width=2.25)
# bottom strip: evaluation discipline
card(s3, 4.95, 6.52, 8.0, 0.36, T_BLUE, line=BORDER, radius=0.08)
textbox(s3, 5.05, 6.52, 7.8, 0.36, "**Evaluation:** hand-verified test set, scored once → citation accuracy "
        f"{EVAL} · refusal accuracy {EVAL}",
        size=9.5, color=NAVY, anchor=MSO_ANCHOR.MIDDLE)
notes(s3, "One Python backend, one PostgreSQL database, one web app. Records flow down from NCPOR's systems; "
          "originals are never copied. The AI only writes from retrieved passages and every sentence is checked.")

# ================================================================== slide 4: feasibility and viability
s4 = S[3]
prepare(s4)
tb = textbox(s4, 0.35, 1.3, 6.5, 0.36, "", size=13.5)
p = tb.text_frame.paragraphs[0]
add_rich(p, "✔ ", 13.5, BLUE, bold=True, font=SYMBOL)
add_rich(p, "FEASIBILITY — already proven, not projected", 13.5, BLUE, bold=True)
stats = [("archive", NAVY, "797", "reports & papers\nharvested, zero entry"),
         ("scan", BLUE, "3,266", "scanned pages\nread & searchable"),
         ("compass", TEAL, "52", "expeditions in\nthe coverage view"),
         ("check", GREEN, "25", "automated tests\npassing")]
for i, (ic, col, big, small) in enumerate(stats):
    x = 0.45 + i * 1.6
    icon_circle(s4, x + 0.3, 1.75, 0.62, ic, col)
    textbox(s4, x, 2.4, 1.3, 0.38, big, size=19, color=col, bold=True, align=PP_ALIGN.CENTER)
    textbox(s4, x - 0.05, 2.8, 1.4, 0.45, small.split("\n"), size=8.5, color=MUTED, align=PP_ALIGN.CENTER, after=0)
card(s4, 0.35, 3.32, 6.45, 0.84, T_PEACH, radius=0.1)
textbox(s4, 0.5, 3.36, 6.2, 0.88, [
    "[[Runs on a laptop CPU:]] search model runs locally — no GPU; one database, one backend",
    "[[Zero new work for NCPOR:]] harvests DSpace + news that already exist; review is the only staff task",
    "[[Open standards, open source:]] OAI-PMH, PostgreSQL, Next.js — zero licence cost",
], size=9.5, bullet="•", bullet_color=BLUE, after=2)

risks = [["⚠  RISKS", "\U0001F6E1  HOW WE HANDLE THEM"],
         ["1980s scans read badly", "Text layer first; OCR confidence shown; original page beside every quote"],
         ["AI states what the source doesn’t", "Per-sentence checker; unsupported blocks approval; refusal threshold"],
         ["Looks like other SIH AI portals", "Real archive, zero entry, coverage view, measured accuracy"],
         ["NPDC not open to harvesting", "Link-only until NCPOR permits; adapter slot ready"],
         ["AI API outage or rate limits", "Retries + cached answers; search works without AI"]]
table(s4, 0.35, 4.3, [2.45, 4.0], 0.41, risks, header_fill=NAVY, font_size=9, align_center_from=2)

tb = textbox(s4, 7.1, 1.3, 5.8, 0.36, "", size=13.5)
p = tb.text_frame.paragraphs[0]
add_rich(p, "➜ ", 13.5, BLUE, bold=True, font=SYMBOL)
add_rich(p, "VIABILITY — path to operations", 13.5, BLUE, bold=True)
cx, cy, rad = 10.0, 4.25, 1.72
ring = shape(s4, MSO_SHAPE.OVAL, cx - rad, cy - rad, rad * 2, rad * 2, WHITE, line=NAVY)
ring.line.width = Pt(3)
hub = shape(s4, MSO_SHAPE.OVAL, cx - 0.85, cy - 0.62, 1.7, 1.24, TEAL)
textbox(s4, cx - 0.82, cy - 0.55, 1.64, 1.1, ["**Polar Index**", "**engine**", "harvest · link", "cite · review"],
        size=9, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, after=0)
phases = [("① Prototype ✓", "Real DSpace + news;\n10 reports read", BLUE, 90),
          ("② NCPOR pilot", "Library / NPDC curators\nconfirm links", NAVY, 0),
          ("③ Full archive", "All reports; 31–45\nfrom digital files", NAVY, 270),
          ("④ Daily operations", "Nightly harvest; feeds\nPolar & Ocean Museum", BLUE, 180)]
for title, body, col, ang in phases:
    a = math.radians(ang)
    px, py = cx + rad * math.cos(a), cy - rad * math.sin(a)
    d = 1.5
    shape(s4, MSO_SHAPE.OVAL, px - d / 2, py - d / 2, d, d, col)
    textbox(s4, px - 0.7, py - 0.55, 1.4, 1.1, [f"**{title}**"] + body.split("\n"), size=8.5, color=WHITE,
            align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, after=0)
notes(s4, "Feasibility numbers are counts from the running prototype. The main external dependency, NPDC access, "
          "is optional: the archive side works on DSpace and the news feed alone.")

# ================================================================== slide 5: impact and benefits
s5 = S[4]
prepare(s5)
ribbon(s5, 0.35, 1.28, 7.4, 0.36, "PRACTICAL IMPACT", fill=BLUE, size=12.5)
textbox(s5, 0.35, 1.7, 7.4, 0.28, "**Case study · a student asks:** “Has India measured how the Dakshin "
        "Gangotri glacier is moving?”", size=10)
shots = [("s5_search.png", "① Search finds the papers", BLUE),
         ("s5_answer.png", "② Cited draft, every sentence checked", ORANGE),
         ("s5_page.png", "③ Click through to the scan", GREEN)]
for i, (img, label, col) in enumerate(shots):
    x = 0.35 + i * 2.5
    textbox(s5, x, 2.0, 2.4, 0.26, label, size=9.5, color=col, bold=True, align=PP_ALIGN.CENTER)
    picture(s5, img, x, 2.28, w=2.4, border=MUTED)
textbox(s5, 0.35, 3.85, 7.4, 0.5, "Answer drawn only from glaciology papers in the 9th, 11th, 26th and 30th "
        "Expedition reports — e.g. the snout ((receded 81 cm in 2011)) — each sentence links to its page. "
        "Off-archive questions are {{refused}}.", size=9, color=DARK)

ribbon(s5, 0.35, 4.42, 3.6, 0.34, "QUANTIFIED IMPACT", fill=NAVY, size=11)
qi = [("6 → 1", "disconnected systems\nbecome one index", BLUE),
      ("0", "new records typed\nby NCPOR staff", GREEN),
      (EVAL, "citation accuracy,\nhand-verified", ORANGE)]
for i, (big, small, col) in enumerate(qi):
    x = 0.35 + i * 1.22
    card(s5, x, 4.85, 1.14, 1.2, T_GREY, radius=0.08)
    textbox(s5, x, 4.88, 1.14, 0.45, big, size=17, color=col, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    textbox(s5, x + 0.03, 5.35, 1.08, 0.65, small.split("\n"), size=8, color=MUTED, align=PP_ALIGN.CENTER, after=0)
ribbon(s5, 4.15, 4.42, 3.6, 0.34, "WHO BENEFITS ?", fill=NAVY, size=11)
users = [("govt", NAVY, "NCPOR outreach\n& MoES"), ("cap", ORANGE, "Teachers &\nstudents"),
         ("microscope", GREEN, "Researchers &\nPhD scholars"), ("newspaper", BLUE, "Journalists\n& media"),
         ("museum", PURPLE, "Polar & Ocean\nMuseum")]
for i, (ic, col, label) in enumerate(users):
    x = 4.15 + (i % 3) * 1.2
    y = 4.85 + (i // 3) * 1.05
    icon_circle(s5, x + 0.38, y, 0.44, ic, col)
    textbox(s5, x, y + 0.46, 1.2, 0.5, label.split("\n"), size=7.5, align=PP_ALIGN.CENTER, after=0)

ribbon(s5, 8.0, 1.28, 4.98, 0.36, "BENEFITS OF THE SOLUTION", fill=NAVY, size=12.5)
benefits = [("people", ORANGE, "SOCIAL", "Students get answers they can trust, cited to India’s own reports"),
            ("rupee", GREEN, "ECONOMIC", "No new data-entry cost; open-source stack, zero licence fees"),
            ("govt", NAVY, "OPERATIONAL", "Outreach posts drafted in minutes from confirmed records"),
            ("chip", PURPLE, "TECHNOLOGICAL", "Every public sentence traceable to a page, with an audit log"),
            ("leaf", TEAL, "ENVIRONMENTAL", "Glacier, ice and climate evidence reaches classrooms"),
            ("flask", BLUE, "SCIENTIFIC", "40 years of expedition science searchable; fewer duplicated baselines"),
            ("globe", RED, "NATIONAL", "India’s polar programme taught with Indian data; Hindi + GIGW 3.0")]
for i, (ic, col, title, body) in enumerate(benefits):
    y = 1.72 + i * 0.73
    shp = shape(s5, MSO_SHAPE.RECTANGLE, 8.0, y, 4.98, 0.3, col)
    icon(s5, ic, 8.08, y + 0.03, 0.24)
    textbox(s5, 8.4, y, 4.5, 0.3, title, size=10, color=WHITE, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    textbox(s5, 8.05, y + 0.31, 4.9, 0.38, body, size=9.5, color=DARK, font="Times New Roman")
notes(s5, "The case study is a real run of the prototype on NCPOR's DSpace archive. The citation-accuracy tile is "
          "filled in from the hand-verified evaluation run.")

# ================================================================== slide 6: research and references
s6 = S[5]
prepare(s6)
caps_label(s6, 0.35, 1.28, 6.3, "Key research findings (built on NCPOR’s real archive)", color=NAVY,
           icon_name="flask", size=12)
findings = [
    ("sitemap2", BLUE, T_BLUE, "The archive exists — the links don’t",
     "Six separate NCPOR systems; DSpace already speaks OAI-PMH, so 32 reports and 765 papers were harvested with "
     "no staff effort"),
    ("scan", ORANGE, T_ORANGE, "Old scans are more readable than feared",
     "9 of 10 reports carry a usable text layer on 68–92% of pages; the 30th was fully OCR’d at 90% mean "
     "confidence"),
    ("listcheck", GREEN, T_GREEN, "AI must be checked sentence by sentence",
     "In our glacier lesson test the checker flagged 4 of 27 checked sentences unsupported — incl. a general "
     "“glaciers move slowly” line absent from the sources"),
]
for i, (ic, col, fill, t, b) in enumerate(findings):
    y = 1.7 + i * 0.98
    card(s6, 0.35, y, 6.3, 0.9, fill, radius=0.08)
    icon_circle(s6, 0.47, y + 0.2, 0.46, ic, col)
    textbox(s6, 1.05, y + 0.05, 5.5, 0.3, t, size=11, color=col, bold=True)
    textbox(s6, 1.05, y + 0.34, 5.5, 0.54, b, size=9, color=DARK)

refs = [
    ("DSpace at NCAOR — Scientific Reports of Indian Expeditions to Antarctica", "http://14.139.119.23:8080/dspace/"),
    ("NCPOR news feed (RSS)", "https://ncpor.res.in/rssfeeds"),
    ("National Polar Data Center (NPDC), NCPOR", "https://npdc.ncpor.res.in/"),
    ("OAI-PMH v2.0 — Open Archives Initiative protocol", "https://www.openarchives.org/OAI/openarchivesprotocol.html"),
    ("Wang et al. (2024) — Multilingual E5 text embeddings, arXiv", "https://arxiv.org/abs/2402.05672"),
    ("Cormack, Clarke & Buettcher (2009) — Reciprocal Rank Fusion, SIGIR", "https://doi.org/10.1145/1571941.1572114"),
    ("Lewis et al. (2020) — Retrieval-Augmented Generation, NeurIPS", "https://arxiv.org/abs/2005.11401"),
    ("WCAG 2.1 — W3C accessibility guidelines (GIGW 3.0 basis)", "https://www.w3.org/TR/WCAG21/"),
]
caps_label(s6, 0.35, 4.66, 6.3, "References :", color=NAVY, size=12)
tb = textbox(s6, 0.45, 4.98, 6.2, 1.6, "", size=9)
tf = tb.text_frame
for i, (label, url) in enumerate(refs):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    p.space_after = Pt(1)
    set_bullet(p, "•", color=BLUE)
    r = p.add_run()
    r.text = label
    r.font.size, r.font.name, r.font.underline = Pt(9), BODY, True
    r.font.color.rgb = BLUE
    r.hyperlink.address = url
textbox(s6, 0.35, 6.62, 6.3, 0.3, ["**github repo link :**   **demo youtube video :**"], size=8.5)

caps_label(s6, 6.95, 1.28, 6.0, "Comparison with existing approaches", color=NAVY, icon_name="chartline", size=12)
Y, N = "✔", "✖"
cmp_rows = [["Capability", "NCPOR\ntoday", "Typical SIH\nprototype", "Polar\nIndex"],
            ["Uses NCPOR’s real archive, no re-entry", N, N, Y],
            ["Links every item to its expedition", N, "partly", Y],
            ["Page-level citation for AI text", N, Y, Y],
            ["Sentence-level fact-check blocks approval", N, N, Y],
            ["Named reviewer + append-only audit", N, "partly", Y],
            ["Measured accuracy on verified set", N, N, Y],
            ["Coverage view of programme gaps", N, N, Y],
            ["Hindi output, reviewed", N, Y, Y]]
hl = {(r, 3): GREEN for r in range(1, len(cmp_rows))}
hl.update({(r, c): RED for r in range(1, len(cmp_rows)) for c in (1, 2) if cmp_rows[r][c] == N})
table(s6, 6.95, 1.7, [3.0, 0.95, 1.1, 0.95], 0.4, cmp_rows, header_fill=NAVY, font_size=9.5, highlight=hl)
textbox(s6, 6.95, 5.38, 6.0, 0.62, "“Typical SIH prototype” = public repositories for PS SIH26063 "
        "reviewed before building (e.g. PolarCROSS, PRISM, POLARIS, DhruvGyani). NCPOR today = main site, DSpace, "
        "NPDC, MET-Data, library and social channels as separate systems.", size=8, color=MUTED, italic=True)
card(s6, 6.95, 6.05, 6.0, 0.8, T_BLUE, line=BORDER, radius=0.08)
textbox(s6, 7.05, 6.08, 5.8, 0.74, [
    "**Coverage view finding:** reports for the 7th and 8th expeditions are listed but unpublished; no DSpace "
    "report exists after the 30th — 12 later expeditions (33–37, 40–46) are known only from NCPOR news."], size=9.5, color=NAVY, anchor=MSO_ANCHOR.MIDDLE)
notes(s6, "Findings are from the running prototype on 30 Sep 2026. References are linked; the comparison is based on "
          "public repositories for this problem statement reviewed before building.")

# ------------------------------------------------------------------ drop the template's instructions slide
ids = prs.slides._sldIdLst
if len(ids) > 6:
    last = list(ids)[6]
    rid = last.get(qn("r:id"))
    ids.remove(last)
    prs.part.drop_rel(rid)

prs.save(OUT)
print("saved", OUT)
